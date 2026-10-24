"""Scoped XLSX/DOCX/PDF jobs. Third-party parsers run in bounded child Jobs."""
from __future__ import annotations
import hashlib
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile
import threading
import time
import zipfile

# The embeddable Windows runtime uses an isolated ._pth and omits the script
# directory. Explicitly add this owned module directory for worker imports.
sys.path.insert(0,str(Path(__file__).resolve().parent))

MAX_DOCUMENT=32*1024*1024
VERSIONS={'openpyxl':'3.1.5','python-docx':'1.2.0','pypdf':'6.10.0','reportlab':'4.4.9'}


def sha(path):
    h=hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda:f.read(1024*1024),b''):h.update(block)
    return h.hexdigest()


def checked_input(files,path,extension):
    p=files.path(path)
    if not p.is_file() or p.suffix.lower()!=extension or p.stat().st_size>MAX_DOCUMENT:
        raise ValueError('Expected '+extension+' file <=32 MiB')
    if extension in ('.xlsx','.docx'):
        with zipfile.ZipFile(p) as z:
            rows=z.infolist()
            if len(rows)>10000 or sum(r.file_size for r in rows)>64*1024*1024 or any(r.flag_bits&1 for r in rows):
                raise ValueError('Encrypted/oversized Office archive refused')
    return p


class Documents:
    def __init__(self,background):self.bg=background;self.lock=threading.Lock()
    def start(self,format,**args):
        from udb_background import identity,publish_new
        from udb_managed_process import ManagedProcess
        actions={'spreadsheet':{'inspect','read','write','edit_cells'},'word_document':{'inspect','read','write','replace_text'},'pdf_document':{'inspect','read','write','merge','extract'}}
        action=args.get('action');extension={'spreadsheet':'.xlsx','word_document':'.docx','pdf_document':'.pdf'}[format]
        if action not in actions[format]:raise ValueError('Unsupported document action')
        files=self.bg.files;given=args.get('path','');path=files.path(given,allow_missing=action in ('write','merge'))
        if path.suffix.lower()!=extension:raise ValueError('Use '+extension+' paths')
        expected=args.get('expected_sha256','NEW');destination=args.get('destination','')
        writing=action in ('write','edit_cells','replace_text','merge','extract')
        inputs=[]
        if action not in ('write','merge'):inputs.append(checked_input(files,given,extension))
        if action=='merge':
            names=args.get('inputs',[])
            if not isinstance(names,list) or not 2<=len(names)<=20:raise ValueError('merge needs 2..20 PDF inputs')
            inputs=[checked_input(files,n,'.pdf') for n in names]
        fingerprints={str(p):sha(p) for p in inputs}
        if action in ('edit_cells','replace_text','extract'):
            if not re.fullmatch('[0-9a-fA-F]{64}',expected) or fingerprints[str(path)]!=expected.lower():
                raise ValueError('Current source expected_sha256 is required for edits/extract')
        elif action in ('read','inspect') and expected!='NEW' and fingerprints[str(path)]!=expected.lower():
            raise ValueError('Source SHA256 mismatch')
        if action=='extract' and not destination:raise ValueError('extract needs a new destination')
        target=files.path(destination,allow_missing=True) if destination else path
        if writing:
            if target.suffix.lower()!=extension or not target.parent.is_dir():raise ValueError('Destination extension/parent is invalid')
            target_expected='NEW' if destination else expected
            if target_expected=='NEW':
                if target.exists():raise FileExistsError('Destination exists; use its SHA256 to replace')
            elif not re.fullmatch('[0-9a-fA-F]{64}',target_expected) or not target.is_file() or target.stat().st_size>MAX_DOCUMENT or sha(target)!=target_expected.lower():
                raise ValueError('Destination SHA256 mismatch')
        else:target_expected='NEW'
        seconds=args.pop('timeout_seconds',30)
        if type(seconds) is not int or not 1<=seconds<=60:raise ValueError('Document timeout 1..60 seconds')
        # Serialize before starting a worker; reject unbounded/invalid input.
        request=dict(format=format,arguments=args,root=str(files.root),inputs={str(p):fingerprints[str(p)] for p in inputs})
        encoded=json.dumps(request,ensure_ascii=False,allow_nan=False).encode('utf-8')
        if len(encoded)>1024*1024:raise ValueError('Document request exceeds 1 MiB')
        def run(task):
            temp=None;proc=None;temporary_identity=None
            try:
                task.deadline=min(task.deadline,time.monotonic()+seconds)
                task.check()
                if writing:
                    fd,name=tempfile.mkstemp(prefix='.udb-stage-office-',suffix=extension,dir=str(target.parent));os.close(fd)
                    temp=Path(name);temporary_identity=identity(temp)[:2]
                config=dict(request,staging=None if temp is None else str(temp))
                proc=ManagedProcess([sys.executable,'-u',str(Path(__file__).resolve()),'--worker'],str(files.root),512)
                task.worker=proc;deadline=time.monotonic()+seconds
                def watchdog():
                    while proc.poll() is None:
                        if task.cancel.wait(.05) or time.monotonic()>deadline:proc.terminate();return
                timer=threading.Thread(target=watchdog,daemon=True);timer.start()
                proc.stdin.write((json.dumps(config,ensure_ascii=False,allow_nan=False)+'\n').encode());proc.stdin.close()
                line=proc.stdout.readline(262145)
                task.check()
                if not line or len(line)>262144:raise RuntimeError('Document worker output missing/too large')
                value=json.loads(line)
                task.check()
                if proc.wait(timeout=3)!=0 or value.get('error'):raise ValueError(value.get('error','Document worker failed'))
                if time.monotonic()>deadline:raise TimeoutError('Document time budget exceeded')
                for source,digest in fingerprints.items():
                    p=checked_input(files,source,extension)
                    if sha(p)!=digest:raise ValueError('Input changed during document operation')
                if writing:
                    if identity(temp)[:2]!=temporary_identity or temp.stat().st_nlink!=1 or temp.stat().st_size>MAX_DOCUMENT:
                        raise ValueError('Document staging identity/size changed')
                    digest=sha(temp)
                    if digest!=value.get('sha256'):raise ValueError('Document output checksum mismatch')
                    with temp.open('ab') as f:f.flush();os.fsync(f.fileno())
                    def commit():
                        with self.lock:
                            p=files.path(str(target),allow_missing=True)
                            if target_expected=='NEW':publish_new(temp,p)
                            else:
                                if not p.is_file() or sha(p)!=target_expected.lower():raise ValueError('Destination changed before commit')
                                recovery=files.root/'.udb-recovery'
                                if recovery.is_symlink() or (recovery.exists() and getattr(recovery.lstat(),'st_file_attributes',0)&0x400):raise ValueError('Recovery directory is redirected')
                                recovery.mkdir(exist_ok=True)
                                import uuid
                                backup=recovery/(uuid.uuid4().hex+'.bak')
                                with p.open('rb') as incoming,backup.open('xb') as out:
                                    import shutil
                                    shutil.copyfileobj(incoming,out,1024*1024);out.flush();os.fsync(out.fileno())
                                if sha(p)!=target_expected.lower():raise ValueError('Destination changed while backing up')
                                os.replace(temp,p)
                    task.commit(commit)
                    value.update(path=str(target.relative_to(files.root)),sha256=digest,committed=True,replay_safe=False)
                return value
            except Exception:
                task.check()  # cancellation/timeout wins over a killed pipe error
                raise
            finally:
                if proc:
                    proc.close();task.worker=None
                    for stream in (proc.stdin,proc.stdout):
                        try:stream.close()
                        except (OSError,ValueError):pass
                if temp is not None and temp.exists() and identity(temp)[:2]==temporary_identity:temp.unlink()
        return self.bg.new(format+'_'+action,run)


def text(value,limit=200000):
    if not isinstance(value,str) or len(value)>limit:raise ValueError('Bounded text required')
    return value


def spreadsheet(files,a,output):
    import openpyxl
    from openpyxl.utils.cell import range_boundaries
    action=a['action'];p=files.path(a['path'],allow_missing=action=='write')
    book=None
    try:
        if action=='write':
            book=openpyxl.Workbook();book.remove(book.active)
            sheets=a.get('sheets',{})
            if not isinstance(sheets,dict) or not 1<=len(sheets)<=10:raise ValueError('Provide 1..10 named sheets')
            count=0
            for name,rows in sheets.items():
                sheet=book.create_sheet(text(name,31))
                if not isinstance(rows,list) or len(rows)>10000:raise ValueError('Bounded row arrays required')
                for row_index,row in enumerate(rows,1):
                    if not isinstance(row,list) or len(row)>100:raise ValueError('Rows have at most 100 cells')
                    for col,value in enumerate(row,1):
                        put_cell(sheet.cell(row_index,col),value,a.get('allow_formulas',False));count+=1
                        if count>10000:raise ValueError('Write limit 10000 cells')
        else:
            book=openpyxl.load_workbook(p,read_only=action in ('read','inspect'),data_only=bool(a.get('data_only',False)) if action=='read' else False,keep_links=action=='edit_cells')
        if action=='inspect':return dict(sheets=[dict(name=s.title,rows=s.max_row,columns=s.max_column) for s in book.worksheets],formulas_evaluated=False)
        if action in ('read','edit_cells'):
            name=a.get('sheet') or book.sheetnames[0]
            if name not in book:raise ValueError('Unknown sheet')
            sheet=book[name]
            if action=='read':
                bounds=range_boundaries(a.get('range','A1:J50'))
                x,y,right,bottom=bounds
                if any(v is None for v in bounds) or x<1 or y<1 or right<x or bottom<y or (right-x+1)*(bottom-y+1)>5000:raise ValueError('Read range limit 5000 cells')
                values=[];cropped=0
                for row in sheet.iter_rows(min_row=y,max_row=bottom,min_col=x,max_col=right,values_only=True):
                    cropped+=sum(isinstance(v,str) and len(v)>2000 for v in row)
                    values.append([v.isoformat() if hasattr(v,'isoformat') else text(v,100000)[:2000] if isinstance(v,str) else v for v in row])
                return dict(sheet=name,range=a.get('range','A1:J50'),rows=values,truncated_cells=cropped,formulas_evaluated=False)
            cells=a.get('cells',{})
            if not isinstance(cells,dict) or not 1<=len(cells)<=10000:raise ValueError('Provide 1..10000 cell edits')
            for address,value in cells.items():
                if not re.fullmatch('[A-Za-z]{1,3}[1-9][0-9]{0,6}',address):raise ValueError('Invalid cell address')
                cell=sheet[address]
                if cell.column>16384 or cell.row>1048576:raise ValueError('Cell outside XLSX limits')
                put_cell(cell,value,a.get('allow_formulas',False))
        book.save(output)
        return dict(sheets=book.sheetnames,formulas_evaluated=False)
    finally:
        if book:book.close()


def put_cell(cell,value,allow_formulas):
    if value is not None and not isinstance(value,(str,bool,int,float)):raise ValueError('Cell values must be text, numbers, boolean or null')
    if isinstance(value,float) and not math.isfinite(value):raise ValueError('Finite numeric cell required')
    if isinstance(value,str):
        text(value,10000);cell.value=value
        if not allow_formulas or not value.startswith('='):cell.data_type='s'
    else:cell.value=value


def paragraphs(document,seen=None):
    if seen is None:seen=set()
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell._tc not in seen:
                    seen.add(cell._tc);yield from paragraphs(cell,seen)


def replace_runs(paragraph,old,new):
    runs=paragraph.runs;value=''.join(r.text for r in runs);count=value.count(old)
    for start in reversed([m.start() for m in re.finditer(re.escape(old),value)]):
        end=start+len(old);offset=0;first=True
        for run in runs:
            source=run.text;right=offset+len(source)
            if right>start and offset<end:
                run.text=source[:max(0,start-offset)]+(new if first else '')+source[max(0,end-offset):]
                first=False
            offset=right
    return count


def word_document(files,a,output):
    from docx import Document
    action=a['action'];doc=Document() if action=='write' else Document(files.path(a['path']))
    if action=='inspect':return dict(paragraphs=len(doc.paragraphs),tables=len(doc.tables))
    if action=='read':
        rows=list(paragraphs(doc));offset=a.get('offset',0);limit=a.get('limit',100)
        if type(offset) is not int or offset<0 or type(limit) is not int or not 1<=limit<=200:raise ValueError('Paragraph offset >=0, limit 1..200')
        return dict(paragraphs=[p.text[:4000] for p in rows[offset:offset+limit]],truncated_paragraphs=[offset+i for i,p in enumerate(rows[offset:offset+limit]) if len(p.text)>4000],next_offset=offset+min(limit,len(rows[offset:])),total=len(rows),scope='body paragraphs and tables')
    if action=='write':
        values=a.get('paragraphs',text(a.get('text','')).splitlines())
        if not isinstance(values,list) or len(values)>2000:raise ValueError('At most 2000 paragraphs')
        for value in values:doc.add_paragraph(text(value,10000))
        tables=a.get('tables',[])
        if not isinstance(tables,list) or len(tables)>20:raise ValueError('At most 20 tables')
        for values in tables:
            if not isinstance(values,list) or not values or len(values)>200:raise ValueError('Table needs 1..200 rows')
            if any(not isinstance(row,list) for row in values):raise ValueError('Table rows must be text arrays')
            columns=max(len(row) for row in values)
            if not 1<=columns<=30:raise ValueError('Table needs 1..30 columns')
            table=doc.add_table(rows=len(values),cols=columns)
            for i,row in enumerate(values):
                for j,value in enumerate(row):table.cell(i,j).text=text(value,4000)
    elif action=='replace_text':
        old=text(a.get('old',''),10000);new=text(a.get('new',''),10000);expected=a.get('occurrences',1)
        if not old or type(expected) is not int or not 1<=expected<=10000:raise ValueError('Nonempty old text and occurrences 1..10000')
        rows=list(paragraphs(doc));actual=sum(''.join(r.text for r in p.runs).count(old) for p in rows)
        if actual!=expected:raise ValueError('Exact occurrence count mismatch')
        for p in rows:replace_runs(p,old,new)
    doc.save(output)
    return dict(paragraphs=len(doc.paragraphs),tables=len(doc.tables),replacements=expected if action=='replace_text' else 0)


def pdf_document(files,a,output):
    from pypdf import PdfReader,PdfWriter
    action=a['action']
    def reader(name):
        value=PdfReader(files.path(name))
        if value.is_encrypted:raise ValueError('Encrypted PDFs are not supported')
        if len(value.pages)>200:raise ValueError('PDF limit 200 pages')
        return value
    if action in ('inspect','read','extract'):
        source=reader(a['path'])
        if action=='inspect':return dict(pages=len(source.pages))
        if action=='read':
            start=a.get('start_page',1);limit=a.get('limit',3)
            if type(start) is not int or not 1<=start<=len(source.pages) or type(limit) is not int or not 1<=limit<=10:raise ValueError('Page range invalid; read at most 10 pages')
            values=[]
            for i in range(start-1,min(len(source.pages),start-1+limit)):
                value=source.pages[i].extract_text() or ''
                values.append(dict(page=i+1,text=value[:16000],text_truncated=len(value)>16000))
            return dict(pages=values,total_pages=len(source.pages),ocr=False)
    if action=='write':
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
        from reportlab.lib.styles import ParagraphStyle
        from xml.sax.saxutils import escape
        candidates=[Path(os.environ.get('SystemRoot','C:/Windows'))/'Fonts'/'arial.ttf',Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')]
        font=next((p for p in candidates if p.is_file()),None)
        if not font:raise ValueError('No installed Unicode TrueType font found')
        pdfmetrics.registerFont(TTFont('UDBUnicode',str(font)))
        style=ParagraphStyle('UDB',fontName='UDBUnicode',fontSize=12,leading=17)
        value=text(a.get('text',''))
        lines=value.splitlines() or ['']
        if len(lines)>10000:raise ValueError('PDF write limit 10000 text paragraphs')
        blocks=[item for line in lines for item in (Paragraph(escape(line) or ' ',style),Spacer(1,5))]
        def page_limit(canvas,document):
            if document.page>200:raise ValueError('PDF write limit 200 pages')
        SimpleDocTemplate(str(output)).build(blocks,onFirstPage=page_limit,onLaterPages=page_limit)
        return dict(pages=len(PdfReader(output).pages),unicode_font=True)
    writer=PdfWriter()
    try:
        if action=='merge':
            pages=0
            for name in a['inputs']:
                source=reader(name);pages+=len(source.pages)
                if pages>200:raise ValueError('Merged PDF limit 200 pages')
                for page in source.pages:writer.add_page(page)
        else:
            selected=a.get('pages',[])
            if not isinstance(selected,list) or not 1<=len(selected)<=200 or any(type(n) is not int or not 1<=n<=len(source.pages) for n in selected):raise ValueError('Provide 1..200 valid 1-based page numbers')
            for n in selected:writer.add_page(source.pages[n-1])
        writer.write(output);return dict(pages=len(writer.pages))
    finally:writer.close()


def worker():
    sys.path.insert(0,str(Path(__file__).resolve().parent/'vendor-office'))
    import logging
    logging.disable(logging.CRITICAL)
    try:
        from udb_companion import ScopedFiles
        line=sys.stdin.buffer.readline(1024*1024+2)
        if len(line)>1024*1024:raise ValueError('Document request too large')
        config=json.loads(line);files=ScopedFiles(config['root']);a=config['arguments']
        extension={'spreadsheet':'.xlsx','word_document':'.docx','pdf_document':'.pdf'}[config['format']]
        for name,digest in config['inputs'].items():
            if sha(checked_input(files,name,extension))!=digest:raise ValueError('Document input changed before parsing')
        output=Path(config['staging']) if config['staging'] else None
        if output and (not output.name.startswith('.udb-stage-office-') or output.is_symlink() or not output.is_relative_to(files.root) or not output.is_file() or output.stat().st_nlink!=1):
            raise ValueError('Invalid issued staging file')
        if output:files.path(str(output.parent))
        # Parsers may warn about unsupported document features. Keep those
        # warnings visible in the result instead of corrupting the JSON pipe.
        import warnings
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter('always')
            value={'spreadsheet':spreadsheet,'word_document':word_document,'pdf_document':pdf_document}[config['format']](files,a,output)
        if caught:
            value['warnings']=[str(item.message)[:300] for item in caught[:10]]
            value['warnings_truncated']=len(caught)>10
        if output:value['sha256']=sha(output)
        value['format']=config['format']
        serialized=json.dumps(value,ensure_ascii=False,allow_nan=False)
        if len(serialized.encode())>250000:raise ValueError('Document result too large; request a smaller range')
        print(serialized,flush=True)
    except Exception as exc:
        print(json.dumps({'error':type(exc).__name__+': '+str(exc)[:300]},ensure_ascii=False),flush=True)
        raise SystemExit(1)


if __name__=='__main__' and sys.argv[1:]==['--worker']:worker()
