"""Accounting workbook review: source immutability, report safety and validation."""

import hashlib
import tempfile
import time
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from udb_companion import Runtime


class InvoiceAuditTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="udb-accounting-")
        self.root = Path(self.temp.name)
        self.rt = Runtime(str(self.root))

    def tearDown(self):
        self.rt.close()
        self.temp.cleanup()

    def sha(self, name):
        return hashlib.sha256((self.root / name).read_bytes()).hexdigest()

    def run_job(self, **kwargs):
        row = self.rt.features.documents.start("spreadsheet", **kwargs)
        end = time.monotonic() + 30
        while row["state"] == "running" and time.monotonic() < end:
            time.sleep(0.02)
            row = self.rt.features.background.task(row["job_id"])
        self.assertEqual(row["state"], "completed", row)
        return row["result"]

    def fixture(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Hoa don"
        sheet.append(["BANG KE HOA DON"])
        sheet.append(["Du lieu kiem tra"])
        sheet.append([
            "Số hóa đơn", "Ký hiệu", "MST người bán",
            "Thành tiền chưa thuế", "Tiền thuế GTGT", "Tổng tiền thanh toán",
        ])
        sheet.append(["0001", "AA/26E", "0123456789", 100000, 10000, 110000])
        sheet.append(["0001", "AA/26E", "0123456789", 100000, 10000, 110000])
        sheet.append(["0002", "AA/26E", "0123456789", 100000, 10000, 120000])
        sheet.append(["", "AA/26E", "0123456789", 50000, 5000, 55000])
        sheet.append(["0004", "AA/26E", "0123456789", "=100000", 10000, 110000])
        sheet.append(["0005", "AA/26E", "0123456789", "1.000.000", "100.000", "1.100.000"])
        book.save(self.root / "input.xlsx")
        book.close()

    def test_audit_detects_review_flags_without_editing_source(self):
        self.fixture()
        before = self.sha("input.xlsx")
        result = self.run_job(action="audit_invoices", path="input.xlsx", sheet="Hoa don")
        self.assertEqual(result["rows_checked"], 6)
        self.assertEqual(result["header_row"], 3)
        self.assertEqual(result["issues_by_code"]["POSSIBLE_DUPLICATE"], 1)
        self.assertEqual(result["issues_by_code"]["AMOUNT_MISMATCH"], 1)
        self.assertEqual(result["issues_by_code"]["MISSING_INVOICE_NO"], 1)
        self.assertEqual(result["issues_by_code"]["INVALID_NET"], 1)
        self.assertEqual(result["amount_totals"]["vat"], "145000")
        self.assertFalse(result["tax_filing_ready"])
        self.assertNotIn("0123456789", str(result))
        self.assertEqual(self.sha("input.xlsx"), before)

    def test_export_report_is_new_and_preserves_source(self):
        self.fixture()
        before = self.sha("input.xlsx")
        result = self.run_job(
            action="export_audit", path="input.xlsx",
            destination="audit.xlsx", expected_sha256=before,
        )
        self.assertTrue(result["committed"])
        self.assertEqual(result["issue_count"], 4)
        self.assertEqual(self.sha("input.xlsx"), before)
        book = load_workbook(self.root / "audit.xlsx")
        self.assertEqual(book.sheetnames, ["Tong hop", "Canh bao"])
        self.assertEqual(book["Canh bao"].max_row, 5)
        self.assertEqual(book["Tong hop"]["A2"].value, "source")
        self.assertEqual(book["Tong hop"]["B2"].data_type, "s")
        book.close()
        with self.assertRaises(FileExistsError):
            self.rt.features.documents.start(
                "spreadsheet", action="export_audit", path="input.xlsx",
                destination="audit.xlsx", expected_sha256=before,
            )

    def test_export_requires_current_hash_and_new_destination(self):
        self.fixture()
        with self.assertRaises(ValueError):
            self.rt.features.documents.start(
                "spreadsheet", action="export_audit", path="input.xlsx",
                destination="report.xlsx",
            )
        with self.assertRaises(ValueError):
            self.rt.features.documents.start(
                "spreadsheet", action="export_audit", path="input.xlsx",
                expected_sha256=self.sha("input.xlsx"),
            )
        with self.assertRaises(ValueError):
            self.rt.features.documents.start(
                "spreadsheet", action="export_audit", path="input.xlsx",
                destination="../escape.xlsx", expected_sha256=self.sha("input.xlsx"),
            )
        self.assertFalse((self.root / "report.xlsx").exists())

    def test_custom_column_mapping_and_locale_amounts(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["Code", "Amount before tax", "Tax", "Paid"])
        sheet.append(["0009", "1.234.567,50", "123.456,75", "1.358.024,25"])
        book.save(self.root / "mapped.xlsx")
        book.close()
        result = self.run_job(
            action="audit_invoices", path="mapped.xlsx",
            column_map={"invoice_no": "A", "net": "B", "vat": "C", "total": "D"},
            tolerance_vnd=0,
        )
        self.assertEqual(result["issue_count"], 0)
        self.assertEqual(result["amount_totals"]["total"], "1358024.25")

    def test_report_treats_untrusted_source_name_as_text(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["Số hóa đơn", "Tổng tiền"])
        sheet.append(["01", 500])
        book.save(self.root / "=2+2.xlsx")
        book.close()
        self.run_job(
            action="export_audit", path="=2+2.xlsx",
            destination="safe.xlsx", expected_sha256=self.sha("=2+2.xlsx"),
        )
        book = load_workbook(self.root / "safe.xlsx")
        self.assertEqual(book["Tong hop"]["B2"].value, "=2+2.xlsx")
        self.assertEqual(book["Tong hop"]["B2"].data_type, "s")
        book.close()


if __name__ == "__main__":
    unittest.main()
