"""Real bounded document workers, round trips, formatting and publication guards."""

import hashlib
import tempfile
import time
import unittest
import zipfile
from pathlib import Path

from udb_companion import Runtime


class DocumentJobs(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="udb-office-")
        self.root = Path(self.temp.name)
        self.rt = Runtime(str(self.root))

    def tearDown(self):
        self.rt.close()
        self.temp.cleanup()

    def run_job(self, name, **args):
        row = self.rt.features.documents.start(name, **args)
        deadline = time.monotonic() + 30
        while row["state"] == "running" and time.monotonic() < deadline:
            time.sleep(0.02)
            row = self.rt.features.background.task(row["job_id"])
        self.assertEqual(row["state"], "completed", row)
        return row["result"]

    def digest(self, name):
        return hashlib.sha256((self.root / name).read_bytes()).hexdigest()

    def test_xlsx_roundtrip_safe_text_and_guarded_edit(self):
        self.run_job(
            "spreadsheet",
            action="write",
            path="data.xlsx",
            sheets={"Dữ liệu": [["Tên", "Số"], ["Tiếng Việt", 3], ["=1+1", 4]]},
        )
        row = self.run_job(
            "spreadsheet",
            action="read",
            path="data.xlsx",
            sheet="Dữ liệu",
            range="A1:B3",
        )
        self.assertEqual(row["rows"], [["Tên", "Số"], ["Tiếng Việt", 3], ["=1+1", 4]])
        from openpyxl import load_workbook

        book = load_workbook(self.root / "data.xlsx")
        self.assertEqual(book.active["A3"].data_type, "s")
        book.close()
        self.run_job(
            "spreadsheet",
            action="edit_cells",
            path="data.xlsx",
            sheet="Dữ liệu",
            cells={"B2": 7},
            expected_sha256=self.digest("data.xlsx"),
        )
        self.assertEqual(
            self.run_job("spreadsheet", action="read", path="data.xlsx", range="B2:B2")[
                "rows"
            ],
            [[7]],
        )
        with self.assertRaises(ValueError):
            self.rt.features.documents.start(
                "spreadsheet",
                action="edit_cells",
                path="data.xlsx",
                cells={"A1": "bad"},
                expected_sha256="0" * 64,
            )
        with self.assertRaises(FileExistsError):
            self.rt.features.documents.start(
                "spreadsheet", action="write", path="data.xlsx", sheets={"x": [[1]]}
            )

    def test_xlsx_styles_and_formulas_survive_targeted_edit(self):
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font

        book = Workbook()
        sheet = book.active
        sheet.title = "Styled"
        sheet["A1"] = "Header"
        sheet["A1"].font = Font(bold=True, color="FF0000")
        sheet.merge_cells("A1:C1")
        sheet["D2"] = "=1+1"
        book.save(self.root / "style.xlsx")
        book.close()
        self.run_job(
            "spreadsheet",
            action="edit_cells",
            path="style.xlsx",
            sheet="Styled",
            cells={"E2": 10},
            expected_sha256=self.digest("style.xlsx"),
        )
        book = load_workbook(self.root / "style.xlsx")
        sheet = book.active
        self.assertTrue(sheet["A1"].font.bold)
        self.assertEqual(sheet["A1"].font.color.rgb, "00FF0000")
        self.assertIn("A1:C1", sheet.merged_cells)
        self.assertEqual(sheet["D2"].value, "=1+1")
        self.assertEqual(sheet["E2"].value, 10)
        book.close()

    def test_docx_split_runs_replacement_preserves_unaffected_format(self):
        from docx import Document

        doc = Document()
        p = doc.add_paragraph()
        p.add_run("Trước ").bold = True
        p.add_run("Xin ").italic = True
        p.add_run("chào")
        p.add_run(" cuối").underline = True
        doc.save(self.root / "source.docx")
        self.run_job(
            "word_document",
            action="replace_text",
            path="source.docx",
            old="Xin chào",
            new="Chào Việt Nam",
            expected_sha256=self.digest("source.docx"),
        )
        doc = Document(self.root / "source.docx")
        p = doc.paragraphs[0]
        self.assertEqual(p.text, "Trước Chào Việt Nam cuối")
        self.assertTrue(p.runs[0].bold)
        self.assertTrue(p.runs[-1].underline)
        self.assertIn(
            "Chào Việt Nam",
            self.run_job("word_document", action="read", path="source.docx")[
                "paragraphs"
            ][0],
        )

    def test_docx_replacement_containing_old_text_changes_each_original_match_once(
        self,
    ):
        self.run_job("word_document", action="write", path="words.docx", text="x x")
        self.run_job(
            "word_document",
            action="replace_text",
            path="words.docx",
            old="x",
            new="xx",
            occurrences=2,
            expected_sha256=self.digest("words.docx"),
        )
        self.assertEqual(
            self.run_job("word_document", action="read", path="words.docx")[
                "paragraphs"
            ],
            ["xx xx"],
        )

    def test_pdf_unicode_merge_extract(self):
        self.run_job(
            "pdf_document",
            action="write",
            path="one.pdf",
            text="Nghiệm thu tiếng Việt\nTổng tiền: 375000 VND",
        )
        row = self.run_job("pdf_document", action="read", path="one.pdf")
        self.assertIn("Nghiệm thu tiếng Việt", row["pages"][0]["text"])
        self.run_job(
            "pdf_document",
            action="merge",
            path="merged.pdf",
            inputs=["one.pdf", "one.pdf"],
        )
        self.assertEqual(
            self.run_job("pdf_document", action="inspect", path="merged.pdf")["pages"],
            2,
        )
        self.run_job(
            "pdf_document",
            action="extract",
            path="merged.pdf",
            destination="page.pdf",
            pages=[2],
            expected_sha256=self.digest("merged.pdf"),
        )
        self.assertEqual(
            self.run_job("pdf_document", action="inspect", path="page.pdf")["pages"], 1
        )

    def test_malformed_document_never_publishes_destination(self):
        with zipfile.ZipFile(self.root / "bad.docx", "w") as z:
            z.writestr("not-a-document", "bad")
        row = self.rt.features.documents.start(
            "word_document",
            action="replace_text",
            path="bad.docx",
            destination="result.docx",
            old="a",
            new="b",
            expected_sha256=self.digest("bad.docx"),
        )
        for _ in range(500):
            row = self.rt.features.background.task(row["job_id"])
            if row["state"] != "running":
                break
            time.sleep(0.01)
        self.assertEqual(row["state"], "error")
        self.assertFalse((self.root / "result.docx").exists())
        self.assertFalse(list(self.root.glob(".udb-stage-office-*")))

    def test_scope_zip_expansion_and_cancellation(self):
        for path in ("../escape.xlsx", "checkpoints/data.xlsx", ".env"):
            with self.assertRaises((ValueError, PermissionError)):
                self.rt.features.documents.start(
                    "spreadsheet", action="write", path=path, sheets={"x": [[1]]}
                )
        with zipfile.ZipFile(self.root / "bomb.docx", "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("large.xml", b" " * (65 * 1024 * 1024))
        with self.assertRaises(ValueError):
            self.rt.features.documents.start(
                "word_document", action="read", path="bomb.docx"
            )
        row = self.rt.features.documents.start(
            "pdf_document", action="write", path="cancelled.pdf", text="hello\n" * 20000
        )
        row = self.rt.features.background.task(row["job_id"], "cancel")
        for _ in range(500):
            row = self.rt.features.background.task(row["job_id"])
            if row["state"] != "running":
                break
            time.sleep(0.01)
        self.assertEqual(row["state"], "cancelled", row)
        self.assertFalse((self.root / "cancelled.pdf").exists())


if __name__ == "__main__":
    unittest.main(verbosity=2)
