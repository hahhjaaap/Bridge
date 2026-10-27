"""Read-only Vietnamese invoice workbook checks and optional separate XLSX report.

This is an accounting *review* tool, not a tax filing or tax calculation engine.
All source paths are checked by the existing scoped document worker.
"""

from __future__ import annotations

import re
import unicodedata
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

MAX_ROWS = 10000
MAX_COLUMNS = 100
MAX_ISSUES = 20000

ALIASES = {
    "invoice_no": ("so hoa don", "so hd", "so hdon", "invoice number"),
    "series": ("ky hieu", "ky hieu hoa don", "ky hieu hd", "invoice series"),
    "seller_tax_id": (
        "mst nguoi ban", "ma so thue nguoi ban", "mst ben ban",
        "ma so thue ben ban", "mst nban", "seller tax id",
    ),
    "date": ("ngay lap", "ngay lap hoa don", "ngay hoa don", "invoice date"),
    "net": (
        "tien truoc thue", "tien chua thue", "thanh tien chua thue",
        "cong tien hang hoa dich vu", "cong tien hang", "doanh so chua thue",
        "tien hang", "net amount",
    ),
    "vat": (
        "tien thue gtgt", "tien thue vat", "thue gtgt",
        "tien thue", "vat amount",
    ),
    "total": (
        "tong tien thanh toan", "tong thanh toan", "tong cong thanh toan",
        "tong tien", "total amount",
    ),
}
NORMALIZED = {field: {_normalize for _normalize in ()} for field in ()}  # populated below


def normalize(value):
    if value is None:
        return ""
    s = unicodedata.normalize("NFKD", str(value).casefold().replace("đ", "d"))
    return re.sub(r"[^a-z0-9]+", " ", "".join(c for c in s if not unicodedata.combining(c))).strip()


NORMALIZED = {field: {normalize(alias) for alias in names} for field, names in ALIASES.items()}


def _column(value, headers):
    if type(value) is int and 1 <= value <= MAX_COLUMNS:
        return value - 1
    if not isinstance(value, str) or len(value) > 100:
        raise ValueError("Column map values must be column letters, index or header")
    from openpyxl.utils.cell import column_index_from_string

    if re.fullmatch(r"[A-Za-z]{1,3}", value):
        try:
            idx = column_index_from_string(value.upper())
            if idx <= MAX_COLUMNS:
                return idx - 1
        except ValueError:
            pass
    hits = [i for i, h in enumerate(headers) if normalize(h) == normalize(value)]
    if len(hits) != 1:
        raise ValueError("Column header not found uniquely: " + value[:60])
    return hits[0]


def _header(sheet, a):
    row_hint = a.get("header_row")
    if row_hint is not None and (type(row_hint) is not int or not 1 <= row_hint <= 25):
        raise ValueError("header_row must be 1..25")
    best = None
    for n, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 25), values_only=True), 1):
        if row_hint is not None and n != row_hint:
            continue
        mapping = {}
        for i, cell in enumerate(row):
            label = normalize(cell)
            for field, names in NORMALIZED.items():
                if label in names and field not in mapping:
                    mapping[field] = i
        score = len(mapping) + (3 if "invoice_no" in mapping else 0)
        if best is None or score > best[0]:
            best = (score, n, list(row), mapping)
    if best is None:
        raise ValueError("No header row found")
    _, header_row, headers, mapping = best
    custom = a.get("column_map", {})
    if not isinstance(custom, dict) or len(custom) > len(ALIASES):
        raise ValueError("column_map must be a small field-to-column object")
    for field, value in custom.items():
        if field not in ALIASES:
            raise ValueError("Unknown accounting field: " + str(field)[:50])
        mapping[field] = _column(value, headers)
    if "invoice_no" not in mapping or not any(k in mapping for k in ("net", "vat", "total")):
        raise ValueError("Map invoice_no and at least one of net, vat or total")
    if len(set(mapping.values())) != len(mapping):
        raise ValueError("Multiple accounting fields mapped to one column")
    return header_row, mapping


def _amount(value):
    if value is None or value == "":
        return None, None
    if isinstance(value, bool):
        return None, "Invalid boolean amount"
    if isinstance(value, (int, float, Decimal)):
        try:
            amount = Decimal(str(value))
            return (amount, None) if amount.is_finite() else (None, "Non-finite amount")
        except InvalidOperation:
            return None, "Invalid numeric amount"
    if not isinstance(value, str) or len(value) > 80:
        return None, "Invalid amount type"
    s = value.strip().replace("\u00a0", "").replace(" ", "")
    s = re.sub(r"(?i)(vnd|vnđ|đ|₫)$", "", s)
    if s.startswith("="):
        return None, "Formula is not evaluated"
    negative = s.startswith("(") and s.endswith(")")
    if negative:
        s = s[1:-1]
    if s.startswith(("+", "-")):
        sign, s = s[0], s[1:]
    else:
        sign = ""
    if not s or not re.fullmatch(r"\d[\d.,]*", s):
        return None, "Unrecognized amount"
    if "," in s and "." in s:
        decimal_sep = "," if s.rfind(",") > s.rfind(".") else "."
        thousands_sep = "." if decimal_sep == "," else ","
        parts = s.split(decimal_sep)
        if len(parts) != 2 or not 1 <= len(parts[1]) <= 2:
            return None, "Ambiguous amount separators"
        whole = parts[0].split(thousands_sep)
        if len(whole) > 1 and (not 1 <= len(whole[0]) <= 3 or any(len(p) != 3 for p in whole[1:])):
            return None, "Invalid thousands grouping"
        s = "".join(whole) + "." + parts[1]
    elif "," in s or "." in s:
        sep = "," if "," in s else "."
        parts = s.split(sep)
        if len(parts) > 1 and all(len(p) == 3 for p in parts[1:]) and 1 <= len(parts[0]) <= 3:
            s = "".join(parts)
        elif len(parts) == 2 and 1 <= len(parts[1]) <= 2:
            s = parts[0] + "." + parts[1]
        else:
            return None, "Ambiguous amount separators"
    try:
        number = Decimal(("-" if negative else sign) + s)
        return (number, None) if number.is_finite() else (None, "Non-finite amount")
    except InvalidOperation:
        return None, "Invalid amount"


def _display(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value[:150]
    return str(value)[:150]


def _put(sheet, row, col, value):
    cell = sheet.cell(row, col)
    if isinstance(value, str):
        cell.value = value[:2000]
        cell.data_type = "s"  # Do not turn untrusted invoice text into formulas.
    else:
        cell.value = value


def analyze(files, a):
    import openpyxl

    source = files.path(a["path"])
    book = openpyxl.load_workbook(source, read_only=True, data_only=False, keep_links=False)
    try:
        sheet_name = a.get("sheet") or book.sheetnames[0]
        if sheet_name not in book:
            raise ValueError("Unknown sheet")
        sheet = book[sheet_name]
        if sheet.max_row > MAX_ROWS + 25 or sheet.max_column > MAX_COLUMNS:
            raise ValueError("Accounting audit limited to 10000 data rows and 100 columns")
        header_row, mapping = _header(sheet, a)
        tolerance = a.get("tolerance_vnd", 1)
        if type(tolerance) is not int or not 0 <= tolerance <= 1000000:
            raise ValueError("tolerance_vnd must be 0..1000000")
        issues = []
        totals = defaultdict(Decimal)
        counts = defaultdict(int)
        keys = {}
        checked = 0

        def issue(row, code, detail):
            if len(issues) >= MAX_ISSUES:
                raise ValueError("Too many issues; split the workbook")
            issues.append((row, code, detail[:250]))

        for row_index, row in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            header_row + 1,
        ):
            if not any(value is not None and value != "" for value in row):
                continue
            checked += 1
            if checked > MAX_ROWS:
                raise ValueError("Accounting audit exceeds 10000 rows")
            values = {key: row[index] if index < len(row) else None for key, index in mapping.items()}
            number = _display(values.get("invoice_no")).strip()
            series = _display(values.get("series")).strip()
            seller = _display(values.get("seller_tax_id")).strip()
            if not number:
                issue(row_index, "MISSING_INVOICE_NO", "Missing invoice number")
            if "seller_tax_id" in mapping and not seller:
                issue(row_index, "MISSING_SELLER_TAX_ID", "Missing seller tax ID")
            if number:
                key = (normalize(seller), normalize(series), normalize(number))
                if key in keys:
                    issue(row_index, "POSSIBLE_DUPLICATE", "Invoice identifier also occurs at row " + str(keys[key]) + "; multiple line items may be legitimate")
                else:
                    keys[key] = row_index
            amounts = {}
            for field in ("net", "vat", "total"):
                if field not in mapping:
                    continue
                value = values.get(field)
                parsed, error = _amount(value)
                if error:
                    issue(row_index, "INVALID_" + field.upper(), error)
                elif parsed is None:
                    issue(row_index, "MISSING_" + field.upper(), "Amount is empty")
                else:
                    totals[field] += parsed
                    counts[field] += 1
                    amounts[field] = parsed
            if all(field in amounts for field in ("net", "vat", "total")):
                delta = amounts["net"] + amounts["vat"] - amounts["total"]
                if abs(delta) > tolerance:
                    issue(row_index, "AMOUNT_MISMATCH", "Net + VAT - total differs by " + str(delta))
        by_code = defaultdict(int)
        for _, code, _ in issues:
            by_code[code] += 1
        return {
            "sheet": sheet_name,
            "header_row": header_row,
            "column_map": {k: v + 1 for k, v in mapping.items()},
            "rows_checked": checked,
            "issue_count": len(issues),
            "issues_by_code": dict(sorted(by_code.items())),
            "amount_totals": {k: str(totals[k]) for k in ("net", "vat", "total") if k in mapping},
            "amount_cells_counted": {k: counts[k] for k in ("net", "vat", "total") if k in mapping},
            "issues_preview": [{"row": n, "code": code} for n, code, _ in issues[:30]],
            "issues": issues,
            "source": source.name,
            "tax_filing_ready": False,
            "note": "Review only; repeated invoice identifiers can be valid line items. No VAT legal classification or formula evaluation.",
        }
    finally:
        book.close()


def accounting_spreadsheet(files, a, output: Path | None):
    report = analyze(files, a)
    issues = report.pop("issues")
    if a["action"] == "audit_invoices":
        return report
    if a["action"] != "export_audit" or output is None:
        raise ValueError("Unsupported accounting action")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    summary = book.active
    summary.title = "Tong hop"
    summary.append(["Chi tieu", "Gia tri"])
    for key in ("source", "sheet", "header_row", "rows_checked", "issue_count"):
        _put(summary, summary.max_row + 1, 1, key)
        _put(summary, summary.max_row, 2, report[key])
    for key, count in report["issues_by_code"].items():
        summary.append(["Issue: " + key, count])
    for key, value in report["amount_totals"].items():
        summary.append(["Sum " + key + " (review only)", value])
    summary.append(["Disclaimer", report["note"]])
    summary.append(["Formula evaluation", "Not performed"])
    summary.append(["Tax filing", "Not prepared"])

    warnings = book.create_sheet("Canh bao")
    warnings.append(["Dong Excel", "Ma canh bao", "Mo ta"])
    for row_number, code, detail in issues:
        n = warnings.max_row + 1
        _put(warnings, n, 1, row_number)
        _put(warnings, n, 2, code)
        _put(warnings, n, 3, detail)
    for ws, widths in ((summary, [35, 90]), (warnings, [18, 28, 90])):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E497D")
    book.save(output)
    book.close()
    return report
